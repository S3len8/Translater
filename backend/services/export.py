from io import BytesIO
from fastapi import HTTPException, status
from backend.repositories.export import ExportRepository
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

class ExportService:
    def __init__(self, repository: ExportRepository):
        self.repository = repository

    async def create_excel(self):
        histories = await self.repository.get_translated_histories()
        if not histories:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No translation history available"
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Translations"

        headers = ["Word", "Translation", "Transcription", "Score"]
        worksheet.append(headers)

        for history in histories:
            worksheet.append([
                history.get("word"),
                history.get("translation"),
                history.get("transcription"),
                history.get("score"),
            ])

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        worksheet.column_dimensions["A"].width = 25
        worksheet.column_dimensions["B"].width = 35
        worksheet.column_dimensions["C"].width = 25
        worksheet.column_dimensions["D"].width = 12

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        excel_file = BytesIO()
        workbook.save(excel_file)
        excel_file.seek(0)

        return excel_file

    async def create_excel_limited(self, start: int, end: int):
        histories = await self.repository.get_limited_histories(start=start, end=end)
        if not histories:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail="No translation history available"
            )

        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Translations"

        headers = ["Word", "Translation", "Transcription", "Score"]
        worksheet.append(headers)

        for history in histories:
            worksheet.append([
                history.get("word"),
                history.get("translation"),
                history.get("transcription"),
                history.get("score"),
            ])

        header_fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )

        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        for row in worksheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

        worksheet.column_dimensions["A"].width = 25
        worksheet.column_dimensions["B"].width = 35
        worksheet.column_dimensions["C"].width = 25
        worksheet.column_dimensions["D"].width = 12

        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        excel_file = BytesIO()
        workbook.save(excel_file)
        excel_file.seek(0)

        return excel_file

